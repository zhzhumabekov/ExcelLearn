"""Заготовка урока из файлов книги: лист — из шаблона (*temp.xlsx), expect — из решения (*final.xlsx).

openpyxl — dev-зависимость: модуль импортируется только командой lesson_import.
"""

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

DEFAULT_COLUMN_WIDTH = 8.43  # ширина колонки Excel по умолчанию, в символах

Bounds = tuple[int, int, int, int]  # min_col, min_row, max_col, max_row


@dataclass
class ImportResult:
    sheet: dict
    expect: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def column_width_px(width: float | None) -> int:
    # Ширина Excel — в цифрах Calibri 11 (~7 px); у плеера шрифт ячеек крупнее (Segoe UI 14 px, цифра ~8 px)
    return round((width or DEFAULT_COLUMN_WIDTH) * 8 + 12)


def cell_format(number_format: str) -> str:
    if "$" in number_format:
        return "usd"
    if "%" in number_format:
        return "percent"
    if re.search(r"0\.0", number_format):
        return "number"
    return "general"


def clean_formula(formula: str) -> str:
    # В файле функции новее Excel 2007 хранятся с префиксом: =_xlfn.FORMULATEXT(E4)
    return formula.replace("_xlfn.", "")


def clean_number(value: float) -> int | float:
    if float(value).is_integer():
        return int(value)
    return round(value, 4)


def _union(a: Bounds, b: Bounds) -> Bounds:
    return min(a[0], b[0]), min(a[1], b[1]), max(a[2], b[2]), max(a[3], b[3])


def _bounds(ws: Worksheet) -> Bounds:
    return range_boundaries(ws.calculate_dimension())


def _cells(ws: Worksheet, bounds: Bounds):
    min_col, min_row, max_col, max_row = bounds
    yield from ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)


def _custom_width(ws: Worksheet, index: int) -> float | None:
    # .get, а не [], — иначе openpyxl создаст колонку с шириной 13 вместо стандартной
    dim = ws.column_dimensions.get(get_column_letter(index))
    return dim.width if dim is not None and dim.customWidth else None


def _is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def collect_expect(final_formulas: Worksheet, final_values: Worksheet, template: Worksheet, bounds: Bounds):
    """Формулы, которых нет в шаблоне, — то, что ученик посчитает в уроке; их значения идут в expect."""
    expect: dict[str, int | float | str] = {}
    missing: list[str] = []
    for row in _cells(final_formulas, bounds):
        for cell in row:
            addr = cell.coordinate
            if not _is_formula(cell.value) or _is_formula(template[addr].value):
                continue
            value = final_values[addr].value
            if value is None:
                missing.append(addr)
            elif isinstance(value, int | float) and not isinstance(value, bool):
                expect[addr] = clean_number(value)
            else:
                expect[addr] = clean_formula(str(value))
    return expect, missing


def build_cells(template: Worksheet, template_values: Worksheet, bounds: Bounds, result_cells: set[str]):
    cells: dict[str, object] = {}
    warnings: list[str] = []
    header_row = None

    rows = list(_cells(template, bounds))
    for row in rows:
        values = [c.value for c in row if c.value is not None]
        if values:
            if all(isinstance(v, str) and not _is_formula(v) for v in values):
                header_row = row[0].row
            break

    for row in rows:
        filled = [c for c in row if c.value is not None]
        for cell in row:
            addr, value = cell.coordinate, cell.value
            fmt = cell_format(cell.number_format)
            spec: dict[str, object] = {}

            if value is None:
                if addr not in result_cells or fmt == "general":
                    continue
            elif _is_formula(value):
                spec["f"] = clean_formula(value)
                cached = template_values[addr].value
                if cached is not None:
                    spec["v"] = clean_number(cached) if isinstance(cached, int | float) else cached
            elif isinstance(value, str):
                text = re.sub(r"\s{2,}", " ", value.strip())
                if text != value:
                    warnings.append(f"{addr}: лишние пробелы убраны: {value!r} → {text!r}")
                spec["v"] = text
                if cell.row == header_row:
                    spec["style"] = "head"
                elif len(filled) == 1:
                    spec["style"] = "label"
            else:
                spec["v"] = clean_number(value) if isinstance(value, int | float) else value

            if fmt != "general":
                spec["fmt"] = fmt
            cells[addr] = spec["v"] if list(spec) == ["v"] else spec
    return cells, warnings


def import_workbook(
    template_path: Path, final_path: Path | None = None, cell_range: str | None = None, sheet_name: str | None = None
) -> ImportResult:
    template_wb = openpyxl.load_workbook(template_path)
    template_values_wb = openpyxl.load_workbook(template_path, data_only=True)
    template = template_wb[sheet_name] if sheet_name else template_wb.active
    template_values = template_values_wb[template.title]

    warnings: list[str] = []
    expect: dict = {}
    if cell_range:
        bounds = range_boundaries(cell_range)
    else:
        bounds = _bounds(template)

    if final_path:
        final_wb = openpyxl.load_workbook(final_path)
        final_values_wb = openpyxl.load_workbook(final_path, data_only=True)
        final = final_wb[template.title] if template.title in final_wb.sheetnames else final_wb.active
        final_values = final_values_wb[final.title]
        if not cell_range:
            bounds = _union(bounds, _bounds(final))
        expect, missing = collect_expect(final, final_values, template, bounds)
        if missing:
            warnings.append("нет сохранённых значений формул (файл не пересчитан в Excel): " + ", ".join(missing))

    cells, cell_warnings = build_cells(template, template_values, bounds, set(expect))
    warnings.extend(cell_warnings)

    max_col, max_row = bounds[2], bounds[3]
    columns = {get_column_letter(i): column_width_px(_custom_width(template, i)) for i in range(1, max_col + 2)}
    sheet = {
        "file": template_path.name,
        "name": template.title,
        "columns": columns,
        "rows": max_row + 1,
        "cells": cells,
    }
    return ImportResult(sheet=sheet, expect=expect, warnings=warnings)

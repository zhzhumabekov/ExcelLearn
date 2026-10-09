"""Файл урока для ученика: лист в том виде, в каком урок начинается, — чтобы повторить шаги в своём Excel.

Обратная операция к xlsx_import: значения, формулы шаблона, форматы чисел, стили заголовков и ширина колонок.
"""

import re
from html import unescape
from io import BytesIO
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .exercises import answer_addresses, solution_sheet
from .schema import CellSpec, Exercise, Lesson, Sheet

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

NUMBER_FORMATS = {
    "usd": '"$"#,##0.00',
    "kzt": '#,##0 "₸"',
    "number": "0.00",
}
HEAD_FILL = PatternFill("solid", fgColor="E8F3EC")  # --xl-green-soft плеера
TOTAL_BORDER = Border(top=Side(style="double", color="999999"))
ANSWER_FILL = PatternFill("solid", fgColor="FFF4CC")  # «впишите сюда формулу»


def column_width_chars(px: int) -> float:
    # Обратно к xlsx_import.column_width_px: px = ширина * 8 + 12
    return max(round((px - 12) / 8, 2), 1.0)


def percent_format(value: object) -> str:
    # Как в плеере: целые проценты — «30%», остальные — с двумя знаками
    if isinstance(value, int | float) and float(round(value * 100, 6)).is_integer():
        return "0%"
    return "0.00%"


def lesson_filename(lesson: Lesson) -> str:
    return lesson.sheet.file or f"{lesson.id}.xlsx"


def build_workbook(sheet: Sheet, *, answers: frozenset[str] = frozenset()) -> Workbook:
    """answers — ячейки, которые заполняет ученик: подсвечены жёлтым (в задании и в решении)."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet.name[:31]  # предел Excel на длину имени листа
    for addr in answers:
        ws[addr].fill = ANSWER_FILL

    for letter, px in sheet.columns.items():
        ws.column_dimensions[letter].width = column_width_chars(px)

    for addr, raw in sheet.cells.items():
        spec = raw if isinstance(raw, CellSpec) else CellSpec(v=raw)
        cell = ws[addr]
        if spec.f is not None:
            cell.value = spec.f
        elif spec.v is not None:
            cell.value = spec.v
            if isinstance(spec.v, str) and spec.v.startswith("="):
                cell.data_type = "s"  # текст, похожий на формулу, остаётся текстом

        if spec.fmt == "percent":
            cell.number_format = percent_format(spec.v)
        elif spec.fmt in NUMBER_FORMATS:
            cell.number_format = NUMBER_FORMATS[spec.fmt]

        if spec.style in ("head", "label", "total"):
            cell.font = Font(bold=True)
        if spec.style == "head":
            cell.fill = HEAD_FILL
        if spec.style == "total":
            cell.border = TOTAL_BORDER

    # Курсор в A1, как при открытии урока
    ws.sheet_view.selection[0].activeCell = "A1"
    ws.sheet_view.selection[0].sqref = "A1"
    return wb


def _save(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def lesson_xlsx(lesson: Lesson) -> bytes:
    return _save(build_workbook(lesson.sheet))


# ---------- Задания: файл задания и файл решения ----------


def plain_text(html: str) -> str:
    """HTML условия → текст для ячейки: <br> и абзацы — переносы строк, остальные теги убираются."""
    text = re.sub(r"<br\s*/?>|</p>", "\n", html)
    text = re.sub(r"<[^>]+>", "", text)
    text = re.sub(r"[ \t]+", " ", unescape(text))
    return "\n".join(line.strip() for line in text.splitlines()).strip()


def exercise_filename(exercise: Exercise, kind: Literal["task", "solution"]) -> str:
    stem = (exercise.sheet.file or f"Задание_{exercise.number}.xlsx").removesuffix(".xlsx")
    return f"{stem}_{'задание' if kind == 'task' else 'решение'}.xlsx"


def _add_info_sheet(wb: Workbook, exercise: Exercise, kind: Literal["task", "solution"]) -> None:
    ws = wb.create_sheet("Условие")
    ws.column_dimensions["A"].width = 100
    hint = (
        "Жёлтые ячейки на первом листе — для ваших формул."
        if kind == "task"
        else "Это решение. Щёлкните по жёлтой ячейке на первом листе — в строке формул видна формула."
    )
    rows = [
        (f"Задание {exercise.number}. {exercise.title}", Font(bold=True, size=14)),
        (hint, Font(italic=True)),
        ("Условие", Font(bold=True)),
        (plain_text(exercise.task), None),
        ("Как решать", Font(bold=True)),
        (plain_text(exercise.explain), None),
    ]
    if exercise.mistake:
        rows += [("Частая ошибка", Font(bold=True)), (plain_text(exercise.mistake), None)]
    for i, (text, font) in enumerate(rows, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            cell.font = font


def exercise_xlsx(exercise: Exercise, kind: Literal["task", "solution"]) -> bytes:
    sheet = exercise.sheet if kind == "task" else solution_sheet(exercise)
    wb = build_workbook(sheet, answers=frozenset(answer_addresses(exercise)))
    _add_info_sheet(wb, exercise, kind)
    return _save(wb)

from io import StringIO
from pathlib import Path

import openpyxl
import pytest
import yaml
from django.conf import settings
from django.core.management import call_command

from lessons.management.commands.lesson_import import render_yaml
from lessons.schema import Sheet
from lessons.xlsx_import import cell_format, clean_formula, collect_expect, import_workbook

USD = r'_("$"* #,##0.00_);_("$"* \(#,##0.00\);_("$"* "-"??_);_(@_)'


@pytest.fixture
def template(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "wages"
    ws["B3"], ws["C3"], ws["D3"], ws["E3"] = "Employee", "Hours", "Wage", "Salary"
    ws["B4"], ws["C4"], ws["D4"] = "Luka  Abrus", 49, 10
    ws["D4"].number_format = USD
    ws["E4"].number_format = USD  # пустая ячейка под результат
    ws["E5"].number_format = USD  # пустая, но в решении её нет — в урок не попадает
    ws["B6"] = "Totals"
    ws["G9"] = "=B4"  # формула уже в шаблоне
    ws.column_dimensions["B"].width = 18.57
    path = tmp_path / "wagestemp.xlsx"
    wb.save(path)
    return path


def test_template_becomes_valid_sheet(template):
    result = import_workbook(template, final_path=None, cell_range=None)
    cells = result.sheet["cells"]

    assert cells["B3"] == {"v": "Employee", "style": "head"}
    assert cells["B4"] == "Luka Abrus"
    assert cells["C4"] == 49
    assert cells["D4"] == {"v": 10, "fmt": "usd"}
    assert cells["B6"] == {"v": "Totals", "style": "label"}
    assert cells["G9"] == {"f": "=B4"}
    assert "E5" not in cells
    assert result.sheet["columns"]["B"] == 161
    assert result.sheet["columns"]["A"] == 79  # ширина по умолчанию
    assert result.sheet["rows"] == 10
    assert any("лишние пробелы" in w for w in result.warnings)
    Sheet.model_validate(result.sheet)


def test_collect_expect_takes_only_new_formulas():
    template = openpyxl.Workbook().active
    template["A1"] = "=1+1"
    formulas = openpyxl.Workbook().active
    values = openpyxl.Workbook().active
    formulas["A1"], values["A1"] = "=1+1", 2  # была в шаблоне — не ожидание
    formulas["E4"], values["E4"] = "=C4*D4", 490.0
    formulas["E14"], values["E14"] = "=AVERAGE(E4:E11)", 449.625
    formulas["F4"], values["F4"] = "=_xlfn.FORMULATEXT(E4)", "=C4*D4"
    formulas["E15"] = "=E14"  # нет сохранённого значения

    expect, missing = collect_expect(formulas, values, template, (1, 1, 6, 15))

    assert expect == {"E4": 490, "F4": "=C4*D4", "E14": 449.625}
    assert missing == ["E15"]


@pytest.mark.parametrize(
    ("number_format", "fmt"),
    [
        ("General", "general"),
        (USD, "usd"),
        ('#,##0 "₸"', "kzt"),
        ("0%", "percent"),
        ("0.000", "number"),
        ("0", "general"),
    ],
)
def test_cell_format(number_format, fmt):
    assert cell_format(number_format) == fmt


def test_clean_formula():
    assert clean_formula("=_xlfn.FORMULATEXT(E29)") == "=FORMULATEXT(E29)"


def test_rendered_yaml_round_trips():
    sheet = {
        "file": "a.xlsx",
        "name": "s",
        "columns": {"A": 64},
        "rows": 5,
        "cells": {"A1": {"v": "Имя", "style": "head"}, "A2": "x: y", "B2": 1.5, "C2": {"f": "=A1"}},
    }
    expect = {"C4": "=SUM(C1:C3)", "E4": 490, "F4": "="}

    data = yaml.safe_load(render_yaml("a.xlsx", sheet, expect))

    assert data == {"sheet": sheet, "expect": expect}
    assert "..." not in render_yaml("a.xlsx", sheet, expect)


def test_command_writes_file(template, tmp_path):
    output = tmp_path / "out.yaml"

    call_command("lesson_import", str(template), "--output", str(output), stdout=StringIO(), stderr=StringIO())

    data = yaml.safe_load(output.read_text(encoding="utf-8"))
    Sheet.model_validate(data["sheet"])


BOOK = Path(settings.BASE_DIR)
WAGES_TEMPLATE = BOOK / "Templates/Chapter 1 Templates/wagestemp.xlsx"
WAGES_FINAL = BOOK / "Practice Files/Chapter 1 Practice Files/wagesfinal.xlsx"


@pytest.mark.skipif(not WAGES_FINAL.exists(), reason="файлы книги не в репозитории")
def test_wages_import_from_book_files():
    # Урок 01-01 теперь на своих данных (тенге), поэтому сверяем импорт с числами рис. 1.1 книги
    result = import_workbook(WAGES_TEMPLATE, WAGES_FINAL)

    def value(cell):
        return cell.get("v") if isinstance(cell, dict) else cell

    cells = {a: value(c) for a, c in result.sheet["cells"].items()}
    assert cells["B4"] == "Luka Abrus"
    assert cells["C4"] == 49
    assert cells["D4"] == 10
    assert result.sheet["cells"]["D4"]["fmt"] == "usd"
    assert result.expect["E4"] == 490
    assert result.expect["C12"] == 320
    assert result.expect["E14"] == pytest.approx(449.625)
    assert result.expect["F4"] == "=C4*D4"

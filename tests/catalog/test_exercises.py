from io import BytesIO
from pathlib import Path
from urllib.parse import unquote

import openpyxl
import pytest
from django.urls import reverse
from pydantic import ValidationError

from lessons.exercises import solution_sheet
from lessons.loader import ContentError
from lessons.registry import Registry, get_registry
from lessons.schema import ExerciseSet
from lessons.xlsx_export import ANSWER_FILL, plain_text

EXERCISES = get_registry().exercises["ch01"]


def workbook(client, exercise_id, kind):
    response = client.get(reverse(f"catalog:exercise-{kind}", args=["ch01", exercise_id]))
    assert response.status_code == 200
    return response, openpyxl.load_workbook(BytesIO(response.content))


def test_page_lists_exercises_with_both_downloads(client):
    html = client.get(reverse("catalog:exercises", args=["ch01"])).content.decode()

    assert html.count("data-exercise=") == len(EXERCISES) == 6
    assert "Квота вылова на озере Алаколь" in html
    assert 'href="/course/ch01/exercises/final-grade/task.xlsx"' in html
    assert 'href="/course/ch01/exercises/final-grade/solution.xlsx"' in html
    assert "Частая ошибка." in html


def test_unknown_chapter_or_exercise_is_404(client):
    assert client.get(reverse("catalog:exercises", args=["ch00"])).status_code == 404  # у главы 0 заданий нет
    assert client.get(reverse("catalog:exercise-task", args=["ch01", "nope"])).status_code == 404


@pytest.mark.parametrize("exercise", EXERCISES, ids=lambda e: e.id)
def test_every_exercise_has_task_and_solution_files(client, exercise):
    for kind, suffix in (("task", "_задание.xlsx"), ("solution", "_решение.xlsx")):
        response, wb = workbook(client, exercise.id, kind)
        assert wb.sheetnames == [exercise.sheet.name, "Условие"]
        filename = unquote(response["Content-Disposition"].split("filename*=utf-8''")[1])
        assert filename == exercise.sheet.file.removesuffix(".xlsx") + suffix


def test_task_file_has_empty_yellow_answer_cells_and_solution_has_formulas(client):
    _, task = workbook(client, "final-grade", "task")
    _, solution = workbook(client, "final-grade", "solution")
    task_ws, solution_ws = task["Оценки"], solution["Оценки"]

    assert task_ws["L10"].value is None
    assert task_ws["L10"].fill.fgColor.rgb == ANSWER_FILL.fgColor.rgb
    # Формула протянута со сдвигом ссылок, $ сохранены
    assert solution_ws["L10"].value == "=$B$1*K10+$B$2*J10"
    assert solution_ws["J10"].value == "=SUM(B10:F10)/SUM(B$5:F$5)"
    assert solution_ws["L10"].number_format == "0.00%"
    assert solution_ws["B2"].value == "=1-B1"
    assert solution_ws["B2"].number_format == "0.00%"  # формат ячейки задания (процент) сохранился
    assert "Это решение" in solution["Условие"]["A2"].value


def test_info_sheet_has_plain_text_task():
    assert (
        plain_text("Формула <code>=C5/(B5/100)^2</code>.<br>Дальше &lt;b&gt;") == "Формула =C5/(B5/100)^2.\nДальше <b>"
    )


def test_solution_sheet_keeps_task_style():
    gini = next(e for e in EXERCISES if e.id == "gini-coefficient")
    cell = solution_sheet(gini).cells["B10"]

    assert (cell.f, cell.fmt, cell.style) == ("=SUM(D4:D8)/($B$1*SUM(B4:B8))", "number", "total")


def test_fill_source_must_be_a_formula():
    with pytest.raises(ValidationError, match="нет формулы C5"):
        ExerciseSet.model_validate(
            {
                "chapter": "ch01",
                "exercises": [
                    {
                        "id": "x",
                        "number": 1,
                        "title": "X",
                        "task": "t",
                        "explain": "e",
                        "sheet": {},
                        "solution": {"cells": {"C5": 1}, "fill": [{"source": "C5", "to": "C5:C9"}]},
                        "expect": {"C5": 1},
                    }
                ],
            }
        )


def test_exercises_of_unknown_chapter_are_rejected():
    registry = get_registry()
    exercise_set = ExerciseSet(chapter="ch99", exercises=list(EXERCISES))

    with pytest.raises(ContentError, match="главы «ch99» нет"):
        Registry(registry.course, [], [(Path("ch99.yaml"), exercise_set)])


def test_course_page_links_exercises_and_last_lesson_invites(client):
    course = client.get(reverse("catalog:course")).content.decode()
    assert 'href="/course/ch01/exercises/"' in course
    assert "6 заданий · файлы Excel" in course

    last = client.get(reverse("catalog:lesson", args=["coffee-price-cost"])).content.decode()
    first = client.get(reverse("catalog:lesson", args=["wages-weekly-salary"])).content.decode()
    assert 'data-testid="chapter-exercises"' in last
    assert 'data-testid="chapter-exercises"' not in first


def test_chapter2_solution_files_contain_names(client):
    def solution(exercise_id):
        response = client.get(reverse("catalog:exercise-solution", args=["ch02", exercise_id]))
        return openpyxl.load_workbook(BytesIO(response.content))

    def task(exercise_id):
        response = client.get(reverse("catalog:exercise-task", args=["ch02", exercise_id]))
        return openpyxl.load_workbook(BytesIO(response.content))

    assert solution("two-warehouses").defined_names["остатки"].attr_text == "Склады!$A$1:$B$3,Склады!$A$6:$B$8"
    assert len(task("two-warehouses").defined_names) == 0  # имена создаёт ученик
    # Относительное имя Excel хранит от A1 с переходом через край: «пять ячеек над B12»
    assert solution("last-five-years").defined_names["последние_5"].attr_text == "Продажи!A1048572:A1048576"


def test_relative_name_stored_like_excel():
    from lessons.xlsx_export import _relative_to_a1

    assert _relative_to_a1("B6", "B7") == "A1048576"  # строка выше
    assert _relative_to_a1("C7", "B7") == "B1"  # ячейка справа
    assert _relative_to_a1("$B$3", "B7") == "$B$3"  # абсолютная часть не сдвигается


def test_placeholders_rendered_on_page_and_in_files(client):
    from lessons.placeholders import localize_formula, render_placeholders

    assert render_placeholders("Функция {fn:VLOOKUP}: {f:=VLOOKUP(A1,B1:C3,2,FALSE)}") == (
        "Функция ВПР: =ВПР(A1;B1:C3;2;ЛОЖЬ)"
    )
    assert localize_formula('=IF(A1<0.5,"x, y",1.5)') == '=ЕСЛИ(A1<0,5;"x, y";1,5)'
    assert render_placeholders("{f:=IF(A1<5,1,2)}") == "=ЕСЛИ(A1&lt;5;1;2)"

    html = client.get(reverse("catalog:exercises", args=["ch03"])).content.decode()
    assert "{fn:" not in html and "ВПР" in html


def test_info_text_starting_with_formula_stays_text(client):
    # Подсказка «<code>=ИНДЕКС(…)</code> …» начинается с «=»: в файле это должен быть текст, а не формула
    response = client.get(reverse("catalog:exercise-task", args=["ch05", "seller-count"]))
    info = openpyxl.load_workbook(BytesIO(response.content))["Условие"]
    cells = [c for row in info.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert cells and all(c.data_type == "s" for c in cells)

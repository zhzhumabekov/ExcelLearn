from io import BytesIO
from urllib.parse import quote

import openpyxl
import pytest
from django.urls import reverse

from lessons.registry import get_registry
from lessons.schema import Sheet
from lessons.xlsx_export import build_workbook, column_width_chars

LESSON_IDS = [lesson.id for lesson in get_registry().lessons]


def download(client, lesson_id):
    return client.get(reverse("catalog:lesson-file", args=[lesson_id]))


@pytest.mark.parametrize("lesson_id", LESSON_IDS)
def test_every_lesson_has_a_downloadable_xlsx(client, lesson_id):
    lesson = get_registry().get(lesson_id)

    response = download(client, lesson_id)

    assert response.status_code == 200
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    assert quote(lesson.sheet.file) in response["Content-Disposition"]  # кириллица — через filename*
    ws = openpyxl.load_workbook(BytesIO(response.content)).active
    assert ws.title == lesson.sheet.name


def test_file_has_lesson_data_formats_and_widths(client):
    ws = openpyxl.load_workbook(BytesIO(download(client, "wages-weekly-salary").content)).active

    assert ws["A2"].value == "Айгерим Сапарова"
    assert ws["B2"].value == 49
    assert ws["C2"].value == 1800
    assert ws["C2"].number_format == '#,##0 "₸"'
    assert ws["D2"].value is None  # результат ученик посчитает сам
    assert ws["D2"].number_format == '#,##0 "₸"'  # но формат уже стоит, как в уроке
    assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("E8F3EC")
    assert ws.column_dimensions["A"].width == column_width_chars(150)


def test_link_under_player(client):
    html = client.get(reverse("catalog:lesson", args=["wages-weekly-salary"])).content.decode()

    assert 'href="/lessons/wages-weekly-salary/file.xlsx"' in html
    assert "Скачать Табель_кафе.xlsx" in html
    assert html.index('id="player"') < html.index("Скачать Табель_кафе.xlsx")


def test_unknown_lesson_file_is_404(client):
    assert download(client, "nope").status_code == 404


def test_formulas_percent_and_text_like_formula():
    sheet = Sheet.model_validate(
        {
            "cells": {
                "A1": {"f": "=SUM(B1:B2)"},
                "A2": {"v": 0.3, "fmt": "percent"},
                "A3": {"fmt": "percent"},
                "A4": {"v": "=C4*D4"},  # текст из FORMULATEXT в шаблоне — не формула
                "A5": {"v": "Итого", "style": "total"},
            }
        }
    )

    ws = build_workbook(sheet).active

    assert ws["A1"].value == "=SUM(B1:B2)" and ws["A1"].data_type == "f"
    assert ws["A2"].number_format == "0%"
    assert ws["A3"].number_format == "0.00%"
    assert ws["A4"].data_type == "s"
    assert ws["A5"].border.top.style == "double"
    assert ws.title == "Лист1"


def test_multi_sheet_lesson_file_has_all_sheets(client):
    wb = openpyxl.load_workbook(BytesIO(download(client, "name-scope").content))

    assert wb.sheetnames == ["Алматы", "Астана", "Свод"]
    assert wb["Свод"]["D4"].value == 15


def test_initial_names_are_written(tmp_path):
    from lessons.schema import Lesson
    from lessons.xlsx_export import lesson_xlsx

    lesson = Lesson.model_validate(
        {
            "id": "x",
            "chapter": "ch02",
            "order": 1,
            "title": "X",
            "sheet": {"name": "Лист 1", "cells": {"B2": 5}},
            "sheets": [{"name": "Свод"}],
            "names": [
                {"name": "ставка", "ref": "B2"},
                {"name": "итог", "ref": "Свод!A1:A3", "scope": "Свод"},
            ],
            "steps": [{"title": "t", "text": "t", "actions": [{"wait": 1}]}],
        }
    )
    wb = openpyxl.load_workbook(BytesIO(lesson_xlsx(lesson)))

    assert wb.defined_names["ставка"].attr_text == "'Лист 1'!$B$2"
    assert wb["Свод"].defined_names["итог"].attr_text == "Свод!$A$1:$A$3"


def test_dates_are_real_excel_dates():
    from datetime import date

    ws = build_workbook(Sheet.model_validate({"cells": {"A1": {"v": "2025-05-01", "fmt": "date"}}})).active

    assert ws["A1"].value == date(2025, 5, 1)
    assert ws["A1"].number_format == "dd.mm.yyyy"
